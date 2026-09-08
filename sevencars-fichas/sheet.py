"""Access to Base_Datos ('Datos' tab): live Google Sheet (gspread, default) or a local xlsx export (fallback)."""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from common import norm_text, normalize_plate, normalize_ref, parse_date, parse_int, parse_number
from gauth import CredentialsMissing, describe_google_error, gspread_client  # noqa: F401  (re-exported)

DEFAULT_SHEET_ID = "1pm2KiO1vXy5Zn7OGe8wjOXhKzUub2QIG5Tjv4GDqEBI"
DEFAULT_SHEET_TAB = "Datos"
DEFAULT_SHEET_GID = 1423934348          # fallback when no tab is called 'Datos'
BASTIDOR_HEADER = "bastidor"
BASTIDOR_MIN_COL = 30                   # AE (0-based index)
COMBUSTIBLE_HEADER = "combustible"
COMBUSTIBLE_MIN_COL = 31                # AF
# Columns the tool may create (header text, first candidate column)
CREATABLE_COLUMNS = {"bastidor": (BASTIDOR_HEADER, BASTIDOR_MIN_COL), "combustible": (COMBUSTIBLE_HEADER, COMBUSTIBLE_MIN_COL)}

# logical field -> header text in row 1 (matched accent/case-insensitively, first occurrence wins)
FIELD_HEADERS = {
    "modelo": "MODELO",
    "matricula": "MATRICULA",
    "fecha_matriculacion": "FECHA MATRICULACION",
    "kms": "kms",
    "motor_cv": "motor cv",
    "cubicaje": "cubicaje",
    "caja": "caja",
    "matriculacion": "matriculacion",
    "matriculacion_num": "matriculacion num",
    "bastidor": BASTIDOR_HEADER,
    "combustible": COMBUSTIBLE_HEADER,
    # publishing (publicar.py)
    "precio_contado": "PRECIO CONTADO",
    "precio_campana": "PRECIO CAMPAÑA",
    "tarifa_financiacion": "TARIFA FINANCIACION",     # J: NORMAL / ESPECIAL / SIN DTO / Consultanos
    "garantia": "GARANTIA",
    "meses_garantia": "MESES GARANTIA FABRICA",
    "url_imagen": "URL IMAGEN",
    "cuota": "cuota",
}
FIELD_LABELS = {"referencia": "referencia", **FIELD_HEADERS}
# Columns that --escribir may fill (only when the target cell is empty). kms is NEVER written: it is informational
# (in the live sheet column X is an ARRAYFORMULA fed from Ventas; a literal there breaks the formula).
WRITABLE_FIELDS = ("matricula", "fecha_matriculacion", "motor_cv", "cubicaje", "caja", "bastidor", "combustible")
NEVER_WRITE_FIELDS = ("kms",)
# Cells publicar.py / --cuotas may fill (empty cells only): AD cuota, G URL IMAGEN
PUBLISH_WRITABLE_FIELDS = ("cuota", "url_imagen")


class SheetError(Exception):
    pass


def col_letter(index: int) -> str:
    """0-based column index -> 'A', 'Z', 'AA'."""
    letters = ""
    n = index + 1
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def col_index(letters: str) -> int:
    """'AE' -> 30 (0-based)."""
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _empty(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_vin(value) -> str:
    if value is None:
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def validate_vin(vin: str) -> bool:
    """17 alphanumeric characters without I, O or Q."""
    vin = normalize_vin(vin)
    return len(vin) == 17 and re.fullmatch(r"[A-HJ-NPR-Z0-9]{17}", vin) is not None


@dataclass
class VehicleRow:
    row_number: int                      # 1-based row in the sheet
    referencia: str
    raw: dict = field(default_factory=dict)      # logical field -> raw cell value
    modelo: str = ""
    matricula: str = ""
    fecha_matriculacion: date | None = None
    kms: int | None = None
    motor_cv: float | None = None
    cubicaje: int | None = None
    caja: str = ""
    matriculacion: str = ""
    matriculacion_num: int | None = None
    bastidor: str = ""
    combustible: str = ""
    precio_contado: float | None = None
    precio_campana: float | None = None
    tarifa_financiacion: str = ""
    garantia: str = ""
    meses_garantia: int | None = None
    url_imagen: str = ""
    cuota: float | None = None

    def is_empty(self, field_name: str) -> bool:
        return _empty(self.raw.get(field_name))


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


class SheetData:
    """Parsed content of the 'Datos' tab, independent of where it came from.
    `formulas` (optional): the same grid rendered as formulas, to know which cells/columns are formula-driven."""

    def __init__(self, values: list[list], source_label: str = "hoja", formulas: list[list] | None = None):
        if not values:
            raise SheetError("La hoja está vacía")
        self.source_label = source_label
        self.header = [("" if h is None else str(h)) for h in values[0]]
        self.formula_cells: set[tuple[int, int]] = set()      # (row_number, col_idx)
        self.formula_columns: set[int] = set()
        for r, frow in enumerate(formulas or [], start=1):
            for c, v in enumerate(frow):
                if _is_formula(v):
                    self.formula_cells.add((r, c))
                    if r >= 2:
                        self.formula_columns.add(c)
        self.columns: dict[str, int] = {"referencia": 0}
        for field_name, header_text in FIELD_HEADERS.items():
            idx = self._find_header(header_text)
            if idx is not None:
                self.columns[field_name] = idx
        self.rows: list[VehicleRow] = []
        for r, row in enumerate(values[1:], start=2):
            ref = normalize_ref(row[0] if row else None)
            if not ref:
                continue
            self.rows.append(self._parse_row(r, ref, row))

    def _find_header(self, header_text: str) -> int | None:
        target = norm_text(header_text)
        for idx, h in enumerate(self.header):
            if norm_text(h) == target:
                return idx
        return None

    def _cell(self, row: list, field_name: str):
        idx = self.columns.get(field_name)
        if idx is None or idx >= len(row):
            return None
        value = row[idx]
        return None if _empty(value) else value

    def _parse_row(self, row_number: int, ref: str, row: list) -> VehicleRow:
        raw = {f: self._cell(row, f) for f in FIELD_HEADERS}
        raw["referencia"] = ref
        modelo = raw.get("modelo")
        return VehicleRow(
            row_number=row_number,
            referencia=ref,
            raw=raw,
            modelo="" if modelo is None else str(modelo).strip(),
            matricula=normalize_plate(raw.get("matricula")),
            fecha_matriculacion=parse_date(raw.get("fecha_matriculacion")),
            kms=parse_int(raw.get("kms")),
            motor_cv=parse_number(raw.get("motor_cv")),
            cubicaje=parse_int(raw.get("cubicaje")),
            caja="" if raw.get("caja") is None else str(raw.get("caja")).strip(),
            matriculacion="" if raw.get("matriculacion") is None else str(raw.get("matriculacion")).strip(),
            matriculacion_num=parse_int(raw.get("matriculacion_num")),
            bastidor=normalize_vin(raw.get("bastidor")),
            combustible="" if raw.get("combustible") is None else str(raw.get("combustible")).strip(),
            precio_contado=parse_number(raw.get("precio_contado")),
            precio_campana=parse_number(raw.get("precio_campana")),
            tarifa_financiacion="" if raw.get("tarifa_financiacion") is None else str(raw.get("tarifa_financiacion")).strip(),
            garantia="" if raw.get("garantia") is None else str(raw.get("garantia")).strip(),
            meses_garantia=parse_int(raw.get("meses_garantia")),
            url_imagen="" if raw.get("url_imagen") is None else str(raw.get("url_imagen")).strip(),
            cuota=parse_number(raw.get("cuota")),
        )

    def column_letter(self, field_name: str) -> str | None:
        idx = self.columns.get(field_name)
        return col_letter(idx) if idx is not None else None

    def column_letters(self) -> dict[str, str]:
        return {f: col_letter(i) for f, i in self.columns.items()}

    def is_formula_column(self, field_name: str) -> bool:
        """True when any data cell of that column holds a formula (e.g. an ARRAYFORMULA in row 2)."""
        idx = self.columns.get(field_name)
        return idx is not None and idx in self.formula_columns

    def cell_has_formula(self, row_number: int, field_name: str) -> bool:
        idx = self.columns.get(field_name)
        return idx is not None and (row_number, idx) in self.formula_cells

    def row_by_number(self, n: int) -> VehicleRow | None:
        return next((r for r in self.rows if r.row_number == n), None)

    @property
    def has_bastidor_column(self) -> bool:
        return "bastidor" in self.columns

    def has_column(self, field_name: str) -> bool:
        return field_name in self.columns

    def header_target(self, field_name: str, exclude: set[str] | None = None) -> tuple[str | None, str]:
        """Where a creatable column ('bastidor' AE.., 'combustible' AF..) is or would be created: (letter, message).
        The first empty header cell at/after the minimum column; a different header there -> refused.
        `exclude`: letters already reserved for other headers in the same batch."""
        if field_name in self.columns:
            return self.column_letter(field_name), "existe"
        if field_name not in CREATABLE_COLUMNS:
            return None, f"la columna '{field_name}' no existe y no se crea automáticamente"
        _, min_col = CREATABLE_COLUMNS[field_name]
        for idx in range(min_col, max(len(self.header), min_col) + 2):
            h = self.header[idx] if idx < len(self.header) else ""
            if _empty(h) and col_letter(idx) not in (exclude or set()):
                beyond = idx >= len(self.header)
                return col_letter(idx), "falta (se crearía" + (", añadiendo la columna a la hoja)" if beyond else ")")
        return None, f"falta y no hay celda de cabecera libre a partir de {col_letter(min_col)}"

    def reserve_header(self, field_name: str, letter: str) -> None:
        """Register (in memory) a header that is about to be created so later targets skip that column."""
        idx = col_index(letter)
        while len(self.header) <= idx:
            self.header.append("")
        self.header[idx] = CREATABLE_COLUMNS.get(field_name, (field_name, 0))[0]
        self.columns[field_name] = idx

    def bastidor_target_column(self) -> tuple[str | None, str]:
        return self.header_target("bastidor")

    def find_by_ref(self, ref: str) -> VehicleRow | None:
        ref = normalize_ref(ref)
        for row in self.rows:
            if row.referencia == ref:
                return row
        return None

    def find_by_plate(self, plate: str) -> VehicleRow | None:
        plate = normalize_plate(plate)
        for row in self.rows:
            if row.matricula and row.matricula == plate:
                return row
        return None

    def duplicates(self) -> tuple[dict[str, list[int]], dict[str, list[int]]]:
        return duplicates(self.rows)

    def rows_with_vin(self, vin: str, except_row: int | None = None) -> list[int]:
        vin = normalize_vin(vin)
        return [r.row_number for r in self.rows if vin and r.bastidor == vin and r.row_number != except_row]


@dataclass
class CellWrite:
    row_number: int
    field_name: str
    column: str
    value: object
    allow_overwrite: bool = False     # only the explicit correction/restore paths may replace a non-empty cell

    @property
    def a1(self) -> str:
        return f"{self.column}{self.row_number}"

    def describe(self) -> str:
        return f"fila {self.row_number}, columna {self.column}: {self.value}"


# ------------------------------------------------------------- data sources
class XlsxSheet:
    """Local export (read only)."""

    can_write = False

    def __init__(self, path: Path | str, tab: str = DEFAULT_SHEET_TAB):
        self.path = Path(path)
        self.tab = tab

    @property
    def label(self) -> str:
        return f"xlsx {self.path.name} / {self.tab}"

    def load(self) -> SheetData:
        import openpyxl
        if not self.path.is_file():
            raise SheetError(f"No existe el archivo xlsx: {self.path}")
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        if self.tab not in wb.sheetnames:
            raise SheetError(f"La pestaña '{self.tab}' no está en {self.path.name} (hay: {', '.join(wb.sheetnames)})")
        values = [list(r) for r in wb[self.tab].iter_rows(values_only=True)]
        wb.close()
        return SheetData(values, self.label)

    def write(self, writes: list[CellWrite]) -> list[CellWrite]:
        raise SheetError("No se escribe en el xlsx local; usá la hoja de Google (sin --sheet) para --escribir.")


class LiveSheet:
    """Base_Datos through gspread (OAuth of editor-fotos-seven or a service account, see gauth.py)."""

    can_write = True

    def __init__(self, sheet_id: str | None = None, tab: str | None = None, allow_browser: bool = True):
        self.sheet_id = sheet_id or os.environ.get("SHEET_ID") or DEFAULT_SHEET_ID
        self.tab = tab or os.environ.get("SHEET_TAB") or DEFAULT_SHEET_TAB
        self.allow_browser = allow_browser
        self._ws = None
        self.last_notes: list[str] = []      # e.g. "se añade la columna AF a la hoja"

    @property
    def label(self) -> str:
        return f"Google Sheet Base_Datos / {self.tab}"

    def _worksheet(self):
        if self._ws is not None:
            return self._ws
        import gspread
        client = gspread_client(self.allow_browser)
        try:
            spreadsheet = client.open_by_key(self.sheet_id)
        except gspread.exceptions.SpreadsheetNotFound:
            raise SheetError("No se encontró la hoja Base_Datos. ¿La cuenta autorizada tiene acceso a la hoja?")
        except Exception as exc:  # network / auth / API disabled
            raise SheetError(f"No se pudo conectar con la hoja Base_Datos: {describe_google_error(exc)}")
        try:
            self._ws = spreadsheet.worksheet(self.tab)
        except gspread.exceptions.WorksheetNotFound:
            try:
                self._ws = spreadsheet.get_worksheet_by_id(DEFAULT_SHEET_GID)
            except Exception:
                raise SheetError(f"La pestaña '{self.tab}' (ni la gid {DEFAULT_SHEET_GID}) existe en la hoja de Google.")
        return self._ws

    def load(self) -> SheetData:
        ws = self._worksheet()
        try:
            values = ws.get_all_values()
            formulas = ws.get_all_values(value_render_option="FORMULA")
        except Exception as exc:
            raise SheetError(f"Error leyendo la hoja de Google: {exc}")
        return SheetData(values, self.label, formulas=formulas)

    def current_values(self, ranges: list[str]) -> list:
        """Current content of single cells (formulas rendered as text), read right before writing. Cells beyond
        the current grid are empty by definition."""
        ws = self._worksheet()
        col_count = int(getattr(ws, "col_count", 0) or 0)
        inside = [r for r in ranges if not col_count or col_index("".join(ch for ch in r if ch.isalpha())) < col_count]
        got: dict[str, str] = {}
        if inside:
            for rng, chunk in zip(inside, ws.batch_get(inside, value_render_option="FORMULA")):
                rows = list(chunk) if chunk else []
                got[rng] = rows[0][0] if rows and rows[0] else ""
        return [got.get(r, "") for r in ranges]

    def write(self, writes: list[CellWrite]) -> list[CellWrite]:
        """One batch update. Guard: a cell that is currently non-empty (value or formula) is refused unless the
        write is explicitly marked allow_overwrite. Returns the refused writes."""
        if not writes:
            return []
        ws = self._worksheet()
        try:
            current = self.current_values([w.a1 for w in writes])
        except Exception as exc:
            raise SheetError(f"No se pudo comprobar el contenido actual de las celdas: {exc}")
        allowed, refused = [], []
        for w, now in zip(writes, current):
            if not _empty(now) and not w.allow_overwrite:
                refused.append(w)
            else:
                allowed.append(w)
        self.last_notes = []
        if allowed:
            # The grid may be narrower than the target column (e.g. AF when the sheet ends at AE): expand it first
            max_idx = max(col_index(w.column) for w in allowed)
            col_count = int(getattr(ws, "col_count", 0) or 0)
            if col_count and max_idx >= col_count:
                needed = max_idx + 1 - col_count
                try:
                    ws.add_cols(needed)
                except Exception as exc:
                    raise SheetError(f"No se pudo ampliar la hoja hasta la columna {col_letter(max_idx)}: {exc}")
                self.last_notes.append(f"se añade{'n' if needed > 1 else ''} {needed} columna{'s' if needed > 1 else ''} "
                                       f"a la hoja (hasta {col_letter(max_idx)})")
            data = [{"range": w.a1, "values": [[w.value]]} for w in allowed]
            try:
                ws.batch_update(data, value_input_option="USER_ENTERED")
            except Exception as exc:
                raise SheetError(f"Error escribiendo en la hoja de Google: {exc}")
        return refused

    def probe(self) -> dict:
        ws = self._worksheet()
        return {"titulo": ws.spreadsheet.title, "pestaña": ws.title, "filas": ws.row_count, "columnas": ws.col_count}


def duplicates(rows: list[VehicleRow]) -> tuple[dict[str, list[int]], dict[str, list[int]]]:
    """(referencias repetidas, matrículas repetidas) as {valor: [filas]} for values present in >1 row."""
    refs: dict[str, list[int]] = {}
    plates: dict[str, list[int]] = {}
    for r in rows:
        if r.referencia:
            refs.setdefault(r.referencia, []).append(r.row_number)
        if r.matricula:
            plates.setdefault(r.matricula, []).append(r.row_number)
    return ({k: v for k, v in refs.items() if len(v) > 1}, {k: v for k, v in plates.items() if len(v) > 1})


def open_sheet(xlsx_path: str | None = None, allow_browser: bool = True):
    """Default: live Google Sheet. With --sheet <xlsx>: local file."""
    if xlsx_path:
        return XlsxSheet(xlsx_path)
    return LiveSheet(allow_browser=allow_browser)


def header_write_for(field_name: str, data: SheetData, exclude: set[str] | None = None) -> CellWrite | None:
    """CellWrite that creates the header of a creatable column if it is missing (None if it exists)."""
    if data.has_column(field_name):
        return None
    letter, msg = data.header_target(field_name, exclude)
    if letter is None:
        raise SheetError(f"No se puede crear la columna '{field_name}': {msg}")
    return CellWrite(row_number=1, field_name=field_name, column=letter, value=CREATABLE_COLUMNS[field_name][0])


def header_write_for_bastidor(data: SheetData) -> CellWrite | None:
    return header_write_for("bastidor", data)
